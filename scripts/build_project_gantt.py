from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


PROJECT_START = date(2026, 3, 23)
ACTIVE_WEEKS = 5
PROJECT_TITLE = "Implementacion PoC Onboarding Cashea"
COMPANY_NAME = "PathPilot x Cashea"
PROJECT_MANAGER = "Victor (PathPilot)"
CHART_TITLE = "CRONOGRAMA GANTT ONBOARDING CASHEA"

SUMMARY_FILL = PatternFill(fill_type="solid", fgColor="FFCCCCCC")
COMPLETED_FILL = PatternFill(fill_type="solid", fgColor="FF5B9BD5")
IN_PROGRESS_FILL = PatternFill(fill_type="solid", fgColor="FFF4B183")
PLANNED_FILL = PatternFill(fill_type="solid", fgColor="FFA9D18E")
DEPENDENCY_FILL = PatternFill(fill_type="solid", fgColor="FF9DC3E6")
DELIVERABLE_FILL = PatternFill(fill_type="solid", fgColor="FFFFC000")
NO_FILL = PatternFill(fill_type=None)


@dataclass(frozen=True)
class Task:
    row: int
    wbs: str
    title: str
    owner: str
    start: date
    due: date
    pct_complete: float
    kind: str = "task"


SUMMARY_ROWS = {
    11: ("1.0", "Semana 1: Implementacion base"),
    16: ("2.0", "Semana 2: Extraccion, normalizacion y validacion"),
    21: ("3.0", "Semana 3: API base disponible"),
    26: ("4.0", "Semana 4: Cobertura extendida y PROD"),
    31: ("5.0", "Semana 5: Estabilizacion y correccion de bugs"),
}


TASKS = [
    Task(12, "1.1", "Implementacion base del repositorio", PROJECT_MANAGER, date(2026, 3, 23), date(2026, 3, 24), 0.0),
    Task(13, "1.2", "GitHub, CI basica y configuracion del stack", PROJECT_MANAGER, date(2026, 3, 23), date(2026, 3, 25), 0.0),
    Task(14, "1.3", "Contrato base de API y workers locales", PROJECT_MANAGER, date(2026, 3, 24), date(2026, 3, 26), 0.0),
    Task(15, "1.4", "Integracion de flujos corriendo localmente", PROJECT_MANAGER, date(2026, 3, 25), date(2026, 3, 27), 0.0),
    Task(17, "2.1", "Extraccion por tipo documental", PROJECT_MANAGER, date(2026, 3, 30), date(2026, 4, 1), 0.0),
    Task(18, "2.2", "Normalizacion y validacion cruzada base", PROJECT_MANAGER, date(2026, 3, 30), date(2026, 4, 2), 0.0),
    Task(19, "2.3", "Integracion con Google Cloud", PROJECT_MANAGER, date(2026, 3, 31), date(2026, 4, 3), 0.0),
    Task(20, "2.4", "Dependencia: Entorno de Staging listo", PROJECT_MANAGER, date(2026, 4, 1), date(2026, 4, 1), 0.0, "dependency"),
    Task(22, "3.1", "API con soporte para RIFs", PROJECT_MANAGER, date(2026, 4, 6), date(2026, 4, 8), 0.0),
    Task(23, "3.2", "API con soporte para Cedulas", PROJECT_MANAGER, date(2026, 4, 6), date(2026, 4, 9), 0.0),
    Task(24, "3.3", "API con soporte para Actas constitutivas", PROJECT_MANAGER, date(2026, 4, 7), date(2026, 4, 10), 0.0),
    Task(25, "3.4", "Entregable: API disponible para RIFs, Cedulas y Actas constitutivas", PROJECT_MANAGER, date(2026, 4, 10), date(2026, 4, 10), 0.0, "deliverable"),
    Task(27, "4.1", "API con soporte para Actas de Asamblea", PROJECT_MANAGER, date(2026, 4, 13), date(2026, 4, 15), 0.0),
    Task(28, "4.2", "API con soporte para Firmas personales", PROJECT_MANAGER, date(2026, 4, 13), date(2026, 4, 16), 0.0),
    Task(29, "4.3", "API con soporte para Emprendimientos", PROJECT_MANAGER, date(2026, 4, 14), date(2026, 4, 17), 0.0),
    Task(30, "4.4", "Dependencia: Entorno de Produccion listo", PROJECT_MANAGER, date(2026, 4, 15), date(2026, 4, 15), 0.0, "dependency"),
    Task(32, "5.1", "Correccion de bugs criticos", PROJECT_MANAGER, date(2026, 4, 20), date(2026, 4, 22), 0.0),
    Task(33, "5.2", "Hardening y regresion funcional", PROJECT_MANAGER, date(2026, 4, 20), date(2026, 4, 23), 0.0),
    Task(34, "5.3", "Ajustes de despliegue y observabilidad", PROJECT_MANAGER, date(2026, 4, 21), date(2026, 4, 24), 0.0),
    Task(35, "5.4", "Cierre de estabilizacion", PROJECT_MANAGER, date(2026, 4, 24), date(2026, 4, 24), 0.0, "deliverable"),
]


def business_days(start: date, count: int) -> list[date]:
    days: list[date] = []
    current = start
    while len(days) < count:
        if current.weekday() < 5:
            days.append(current)
        current += timedelta(days=1)
    return days


def week_label(start: date) -> str:
    end = start + timedelta(days=4)
    return f"{start.strftime('%b')} {start.day}-{end.strftime('%b')} {end.day}"


def clear_task_row(ws, row: int) -> None:
    for col in range(2, 9):
        ws.cell(row, col).value = None
    for col in range(9, 69):
        ws.cell(row, col).fill = NO_FILL


def apply_summary_row(ws, row: int, wbs: str, title: str, active_end_col: int) -> None:
    clear_task_row(ws, row)
    ws.cell(row, 2).value = wbs
    ws.cell(row, 3).value = title
    ws.cell(row, 2).fill = SUMMARY_FILL
    ws.cell(row, 3).fill = SUMMARY_FILL
    for col in range(9, active_end_col + 1):
        ws.cell(row, col).fill = SUMMARY_FILL


def task_fill(task: Task) -> PatternFill:
    if task.kind == "dependency":
        return DEPENDENCY_FILL
    if task.kind == "deliverable":
        return DELIVERABLE_FILL
    if task.pct_complete >= 1.0:
        return COMPLETED_FILL
    if task.pct_complete > 0:
        return IN_PROGRESS_FILL
    return PLANNED_FILL


def apply_task_row(ws, task: Task, day_to_col: dict[date, int]) -> None:
    clear_task_row(ws, task.row)
    ws.cell(task.row, 2).value = task.wbs
    ws.cell(task.row, 3).value = task.title
    ws.cell(task.row, 4).value = task.owner
    ws.cell(task.row, 5).value = task.start
    ws.cell(task.row, 6).value = task.due
    ws.cell(task.row, 7).value = f"=DAYS360(E{task.row},F{task.row})"
    ws.cell(task.row, 8).value = task.pct_complete

    fill = task_fill(task)
    current = task.start
    while current <= task.due:
        col = day_to_col.get(current)
        if col is not None:
            ws.cell(task.row, col).fill = fill
        current += timedelta(days=1)


def update_headers(ws, workdays: list[date], active_end_col: int) -> None:
    ws["B2"] = CHART_TITLE
    ws["B4"] = "TITULO DEL PROYECTO"
    ws["D4"] = PROJECT_TITLE
    ws["I4"] = "EMPRESA"
    ws["P4"] = COMPANY_NAME
    ws["B5"] = "RESPONSABLE"
    ws["D5"] = PROJECT_MANAGER
    ws["I5"] = "FECHA"
    ws["P5"] = date.today()

    phase_headers = {
        "I8": "Plan semanal de implementacion",
        "X8": "Dependencias y disponibilidad de API",
        "AM8": "",
        "BB8": "",
    }
    for cell, value in phase_headers.items():
        ws[cell] = value

    week_starts = [workdays[index] for index in range(0, len(workdays), 5)]
    week_label_cells = ["I9", "N9", "S9", "X9", "AC9", "AH9", "AM9", "AR9", "AW9", "BB9", "BG9", "BL9"]
    for index, cell in enumerate(week_label_cells):
        ws[cell] = week_label(week_starts[index]) if index < len(week_starts) else None

    day_letters = ["L", "M", "X", "J", "V"]
    for offset, col in enumerate(range(9, 69)):
        cell = ws.cell(10, col)
        if col <= active_end_col:
            cell.value = day_letters[offset % 5]
        else:
            cell.value = None

    for col in range(active_end_col + 1, 69):
        for row in (8, 9, 10):
            ws.cell(row, col).fill = NO_FILL


def build_gantt(template_path: Path, output_path: Path) -> None:
    workdays = business_days(PROJECT_START, ACTIVE_WEEKS * 5)
    day_to_col = {day: 9 + offset for offset, day in enumerate(workdays)}
    active_end_col = 8 + len(workdays)

    workbook = load_workbook(template_path)
    sheet = workbook["Gantt Chart"]

    update_headers(sheet, workdays, active_end_col)

    for row, summary in SUMMARY_ROWS.items():
        apply_summary_row(sheet, row, summary[0], summary[1], active_end_col)

    for task in TASKS:
        apply_task_row(sheet, task, day_to_col)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a project-specific Gantt workbook from the generic template."
    )
    parser.add_argument(
        "--template",
        default="/Users/vitupro14/Downloads/Gantt chart.xlsx",
        type=Path,
        help="Path to the generic Gantt template workbook.",
    )
    parser.add_argument(
        "--output",
        default="docs/Cashea_Onboarding_Agent_Gantt.xlsx",
        type=Path,
        help="Path where the customized workbook will be written.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_gantt(args.template, args.output)


if __name__ == "__main__":
    main()
